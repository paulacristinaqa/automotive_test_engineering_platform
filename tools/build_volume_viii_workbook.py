from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "workbook-volume-viii.md"
OUTPUT = ROOT / "docs" / "ATEP_Volume_VIII_Test_Framework_Engineering_Workbook.docx"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def borders(cell) -> None:  # type: ignore[no-untyped-def]
    properties = cell._tc.get_or_add_tcPr()
    table_borders = properties.first_child_found_in("w:tcBorders")
    if table_borders is None:
        table_borders = OxmlElement("w:tcBorders")
        properties.append(table_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        table_borders.append(element)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, (cell, value) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.width = Inches(widths[index])
        shade(cell, "17365D")
        borders(cell)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values, strict=True)):
            cell.width = Inches(widths[index])
            borders(cell)
            if row_index % 2:
                shade(cell, "EAF0F7")
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def bullets(doc: Document, values: list[str]) -> None:
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ATEP Volume VIII Test Framework Engineering Workbook")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Version 0.7.0   Complete Volume VIII Baseline").bold = True
    doc.add_paragraph(
        "This workbook records the catalog, execution-binding, scheduled-selection, bounded "
        "performance evidence, fault-campaign, and mutation-quality baseline "
        "for ATEP. Reusable "
        "definitions describe test intent, suites preserve reviewed composition, and "
        "catalog-backed "
        "runs now persist deterministic case results and aggregate outcomes."
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Status", "VIII-1 through VIII-7 complete"],
            ["Technology", "FastAPI, PostgreSQL, SQLAlchemy, Alembic, Pydantic"],
            ["Security", "test_catalog:read and test_catalog:manage"],
            ["Cost", "Local first with no paid cloud, AI, or GPU dependency"],
            ["Next", "Volume IX AI Test Engineer increments"],
        ],
        [1.5, 5.4],
    )

    doc.add_heading("1 Scope and Outcome", level=1)
    doc.add_paragraph(
        "VIII-1 separates reusable test design from test execution. A definition captures one "
        "test's preconditions, actions, inputs, expectations, classification, and resource budget. "
        "A suite selects active definitions in an explicit order and freezes their versions. "
        "VIII-2 binds that snapshot to a run and records each case outcome. VIII-3 schedules "
        "that immutable intent for later smoke, sanity, or regression execution. VIII-4 adds "
        "bounded performance and stress evidence with historical comparison. VIII-5 adds "
        "cross-domain fault campaigns with recovery plans and evidence. VIII-6 adds bounded "
        "mutation scoring and explicit requirement coverage gaps."
    )
    bullets(
        doc,
        [
            (
                "Definitions cover core, vehicle, ECU, CAN, diagnostics, EV, ADAS, and "
                "integration domains."
            ),
            "Suites classify smoke, sanity, regression, performance, stress, and safety intent.",
            (
                "Performance and stress now use bounded profiles, deterministic thresholds, "
                "and comparable evidence; the complete Volume VIII baseline is implemented."
            ),
        ],
    )

    doc.add_page_break()
    doc.add_heading("2 Architecture", level=1)
    add_table(
        doc,
        ["Layer", "Responsibility"],
        [
            ["FastAPI", "Bounded contracts, pagination, status filtering, and RBAC"],
            ["Domain service", "Idempotency, lifecycle, snapshots, audit, and outbox"],
            ["PostgreSQL", "Definitions, suite composition snapshots, versions, and status"],
            ["Alembic", "Linear migrations 0051 through 0058 with reversible schema changes"],
            [
                "Run executor",
                "Materializes and updates ordered cases without rewriting catalog intent",
            ],
            [
                "Fault campaigns",
                "Snapshots bounded injection and recovery intent without arbitrary commands",
            ],
            [
                "Performance evidence",
                "Stores bounded profiles, deterministic thresholds, and historical comparisons",
            ],
        ],
        [1.55, 5.35],
    )

    doc.add_heading("3 Definition Contract", level=1)
    add_table(
        doc,
        ["Concept", "Control", "Engineering purpose"],
        [
            ["Identity", "Lowercase URL-safe ID", "Stable API and evidence references"],
            ["Classification", "Domain, level, automation mode", "Selection and reporting"],
            ["Budget", "Timeout from 1 to 86400 seconds", "Bound future executor resources"],
            ["Preconditions", "At most twenty statements", "Make setup assumptions explicit"],
            ["Steps", "One to one hundred unique structured steps", "Preserve executable intent"],
            ["Inputs", "JSON limited to 8192 bytes per step", "Prevent unbounded payloads"],
        ],
        [1.3, 2.45, 3.15],
    )

    doc.add_page_break()
    doc.add_heading("4 Suite Composition", level=1)
    doc.add_paragraph(
        "A suite contains one to two hundred unique active definitions. Every case has a unique "
        "order, required flag, and bounded parameter overrides. The stored composition adds the "
        "definition name and version, creating a deterministic snapshot for future execution."
    )
    add_table(
        doc,
        ["Rule", "Rejected condition", "Reason"],
        [
            ["Known reference", "Definition does not exist", "Prevent dangling suites"],
            ["Active reference", "Definition is draft or archived", "Use reviewed test intent"],
            ["Unique definition", "Same definition appears twice", "Prevent accidental repetition"],
            ["Unique order", "Two cases share an order", "Guarantee deterministic sequence"],
            ["Version snapshot", "None", "Protect intent from later lifecycle changes"],
        ],
        [1.45, 2.55, 2.9],
    )

    doc.add_heading("5 Lifecycle and Concurrency", level=1)
    doc.add_paragraph(
        "Definitions and suites move only from draft to active to archived. A client supplies the "
        "expected version for every transition. Stale versions and backward transitions return "
        "stable HTTP 409 errors before mutation."
    )

    doc.add_heading("6 API and Access Control", level=1)
    add_table(
        doc,
        ["Resource", "Operations", "Permission"],
        [
            ["Test definitions", "Create and status change", "test_catalog:manage"],
            ["Test definitions", "List and detail", "test_catalog:read"],
            ["Test suites", "Create and status change", "test_catalog:manage"],
            ["Test suites", "List and detail", "test_catalog:read"],
            ["Fault campaigns", "Create, status, execute, cancel", "test_catalog:manage"],
            ["Fault campaigns", "List, detail, and results", "test_catalog:read"],
            ["Mutation and coverage", "Create, status, execute, update", "test_catalog:manage"],
            ["Mutation and coverage", "List, detail, and results", "test_catalog:read"],
        ],
        [1.65, 3.25, 2.0],
    )

    doc.add_page_break()
    doc.add_heading("7 Execution Binding", level=1)
    doc.add_paragraph(
        "A run may reference an active catalog suite. Creation freezes the suite identifier, name, "
        "type, version, and composition and creates one pending result per ordered case in the "
        "same "
        "transaction. Existing runs remain reproducible after the source suite is archived."
    )
    add_table(
        doc,
        ["Persisted input", "Bound", "Purpose"],
        [
            ["Attempt", "0 to 100", "Represent controlled execution attempts"],
            ["Duration", "0 to 86400000 ms", "Record terminal execution time"],
            ["Observed result", "4000 characters", "Explain the measured outcome"],
            ["Evidence", "20 references", "Link artifacts without copying content"],
            ["Case page", "200 records", "Protect query and response resources"],
        ],
        [1.65, 1.8, 3.55],
    )

    doc.add_heading("8 Result Aggregation", level=1)
    add_table(
        doc,
        ["Condition", "Run outcome", "Reason"],
        [
            ["Cases remain pending or running", "Running", "Execution is incomplete"],
            ["Required case failed", "Failed", "Required intent was not satisfied"],
            ["Required case skipped", "Failed", "Required evidence is missing"],
            ["All required cases passed", "Passed", "Reviewed acceptance intent was met"],
            [
                "Optional case failed or skipped",
                "Passed if required cases pass",
                "Optional scope is informative",
            ],
        ],
        [2.25, 2.1, 2.65],
    )
    doc.add_paragraph(
        "Case results exclusively derive running, passed, and failed states for catalog-backed "
        "runs. An authorized operator may still cancel an incomplete run directly."
    )

    doc.add_page_break()
    doc.add_heading("9 Scheduler and Selection", level=1)
    doc.add_paragraph(
        "A catalog-backed job requires a suite identifier and matching smoke, sanity, or "
        "regression policy. Scheduling freezes suite identity, version, type, tags, and ordered "
        "composition. "
        "Dispatch later creates the run and pending cases from that stored snapshot without "
        "re-reading mutable catalog data. Legacy jobs remain supported."
    )
    add_table(
        doc,
        ["Selection rule", "Accepted value", "Engineering purpose"],
        [
            ["Policy", "Smoke, sanity, regression", "Keep VIII-3 scope explicit"],
            ["Type match", "Run, policy, and suite agree", "Prevent misleading execution intent"],
            ["Lifecycle", "Active suite at scheduling", "Select reviewed content"],
            ["Snapshot", "Versioned ordered cases", "Make delayed execution reproducible"],
            ["Dispatch", "Stored snapshot only", "Survive later catalog archival"],
        ],
        [1.55, 2.25, 3.1],
    )

    doc.add_page_break()
    doc.add_heading("10 Performance and Stress", level=1)
    doc.add_paragraph(
        "VIII-4 stores bounded workload profiles and immutable measured evidence without running a "
        "load generator in FastAPI. Every execution links to a terminal test run and passes only "
        "when all configured minimum and maximum thresholds pass."
    )
    add_table(
        doc,
        ["Control", "Bound", "Purpose"],
        [
            ["Duration", "3600 seconds total", "Bound execution exposure"],
            ["Virtual users", "500 per stage", "Prevent unsafe concurrency"],
            ["Request rate", "1000 per second", "Protect shared services"],
            ["CPU and memory", "4 cores and 4096 MB", "Keep resource budgets explicit"],
            ["GPU", "Always disabled", "Protect workstation resources"],
            ["Baseline", "Same profile only", "Make trends comparable"],
        ],
        [1.6, 2.0, 3.3],
    )

    doc.add_page_break()
    doc.add_heading("11 Fault Campaigns", level=1)
    doc.add_paragraph(
        "VIII-5 defines reusable campaigns across Digital Vehicle, ECU, CAN, diagnostics, "
        "Electric Vehicle, and ADAS. Each step selects a domain-specific allowlisted action, "
        "declares an expected effect, and includes a mandatory recovery plan. The framework owns "
        "intent and evidence; native simulators retain ownership of the actual state mutation."
    )
    add_table(
        doc,
        ["Safety control", "Bound", "Engineering purpose"],
        [
            ["Blast radius", "Component, network, or vehicle", "Exclude fleet-wide impact"],
            ["Campaign size", "1 to 32 steps", "Bound orchestration and persistence"],
            ["Parameters", "8192 bytes per step", "Prevent unbounded structured input"],
            ["Time budget", "30 minutes total", "Bound injection and recovery exposure"],
            ["Recovery", "Required for every step", "Make restoration and proof explicit"],
            ["Command surface", "Allowlisted actions only", "Exclude arbitrary command execution"],
        ],
        [1.5, 2.35, 3.05],
    )
    doc.add_paragraph(
        "An active campaign creates an immutable execution snapshot and one pending result per "
        "ordered step. Results move through injection, observation, recovery, and terminal states "
        "using optimistic versions. Required failure or skip fails the completed execution; "
        "optional failures remain informative. Exact replay remains valid after campaign archival."
    )

    doc.add_page_break()
    doc.add_heading("12 Mutation Testing and Coverage", level=1)
    doc.add_paragraph(
        "VIII-6 binds an active catalog suite to a reviewed mutation campaign. Each campaign "
        "preserves the suite snapshot and contains up to five hundred uniquely ordered mutants "
        "using seven explicit operators. Execution materializes pending results atomically and "
        "never exposes arbitrary source or shell execution through the API."
    )
    add_table(
        doc,
        ["Control", "Rule", "Engineering purpose"],
        [
            ["Operators", "Seven allowlisted values", "Keep adapters portable and reviewable"],
            ["Parameters", "8192 bytes per mutant", "Bound structured orchestration intent"],
            ["Killed result", "At least one detecting test", "Make detection evidence explicit"],
            [
                "Mutation score",
                "Killed divided by killed plus survived",
                "Avoid distortion by errors",
            ],
            [
                "Required outcome",
                "Survived, error, or skip fails",
                "Protect reviewed acceptance intent",
            ],
            ["Evidence", "At most 20 references", "Link artifacts without copying content"],
        ],
        [1.5, 2.5, 2.9],
    )
    doc.add_paragraph(
        "Requirement records reference known catalog definitions and external evidence. Both "
        "present means covered, only one present means partial, and neither means gap. Collection "
        "responses include totals for each state so dashboards and quality gates can expose "
        "missing tests or missing evidence directly."
    )

    doc.add_page_break()
    doc.add_heading("13 Cross Platform Automation Reporting", level=1)
    doc.add_paragraph(
        "VIII-7 closes the functional baseline with one immutable report per terminal test run. "
        "The report correlates the authoritative vehicle and run with Vehicle Gateway telemetry "
        "and commands, optional fault and mutation executions, and the status and version that "
        "CarSystemUI displayed. Existing REST, WebSocket, workload identity, lease, and evidence "
        "contracts remain authoritative."
    )
    add_table(
        doc,
        ["Evidence", "Consistency rule", "Bound"],
        [
            ["Test run", "Terminal and owned by the vehicle", "One report per run"],
            ["Gateway", "Telemetry publisher and command consumer", "One selected module"],
            ["Telemetry", "Same vehicle and source gateway", "100 identifiers"],
            ["Commands", "Terminal and same vehicle, gateway, and run", "100 identifiers"],
            ["Fault and mutation", "Terminal and same vehicle and run", "One execution each"],
            ["CarSystemUI", "Displayed status and version match the run", "50 observations"],
        ],
        [1.6, 3.7, 1.6],
    )
    doc.add_paragraph(
        "A failed correlated component makes the report failed. If nothing failed but one "
        "component "
        "was cancelled, the report is cancelled; otherwise it passes. Exact retries return the "
        "original report. Audit and outbox evidence retains identities, outcome, and counts "
        "without "
        "copying raw telemetry, commands, or UI observations."
    )

    doc.add_page_break()
    doc.add_heading("14 Events and Audit", level=1)
    add_table(
        doc,
        ["Mutation", "Outbox event", "Audit action"],
        [
            ["Create definition", "atep.test_definition.created.v1", "test_definition.created"],
            [
                "Change definition status",
                "atep.test_definition.status_changed.v1",
                "test_definition.status_changed",
            ],
            ["Create suite", "atep.test_suite.created.v1", "test_suite.created"],
            [
                "Change suite status",
                "atep.test_suite.status_changed.v1",
                "test_suite.status_changed",
            ],
            [
                "Record case result",
                "atep.test_case.result_recorded.v1",
                "test_case.result_recorded",
            ],
            ["Schedule selection", "atep.test_job.scheduled.v1", "test_job.scheduled"],
            ["Dispatch selection", "atep.test_job.dispatched.v1", "test_job.dispatched"],
            ["Create fault campaign", "atep.fault_campaign.created.v1", "fault_campaign.created"],
            [
                "Change campaign status",
                "atep.fault_campaign.status_changed.v1",
                "fault_campaign.status_changed",
            ],
            [
                "Create fault execution",
                "atep.fault_campaign.execution.requested.v1",
                "fault_campaign.execution_requested",
            ],
            [
                "Record fault step",
                "atep.fault_campaign.step_recorded.v1",
                "fault_campaign.step_recorded",
            ],
            [
                "Cancel fault execution",
                "atep.fault_campaign.execution.cancelled.v1",
                "fault_campaign.execution_cancelled",
            ],
            [
                "Create mutation campaign",
                "atep.mutation_campaign.created.v1",
                "mutation_campaign.created",
            ],
            [
                "Create mutation execution",
                "atep.mutation_execution.created.v1",
                "mutation_execution.created",
            ],
            ["Record mutant result", "atep.mutant_result.recorded.v1", "mutant_result.recorded"],
            [
                "Update requirement",
                "atep.requirement_coverage.updated.v1",
                "requirement_coverage.updated",
            ],
            [
                "Create performance profile",
                "atep.performance.profile.created.v1",
                "performance.profile_created",
            ],
            [
                "Record performance execution",
                "atep.performance.execution.recorded.v1",
                "performance.execution_recorded",
            ],
            [
                "Create automation report",
                "atep.cross_platform_automation.report.created.v1",
                "cross_platform_automation.report_created",
            ],
        ],
        [1.65, 3.15, 2.1],
    )
    doc.add_paragraph(
        "The resource, audit record, and outbox event share one database transaction. Exact create "
        "replays return the original resource without duplicate evidence; changed reuse of the "
        "same "
        "public identifier returns a stable conflict."
    )

    doc.add_page_break()
    doc.add_heading("15 Engineering Decisions", level=1)
    add_table(
        doc,
        ["Decision", "Rationale", "Consequence"],
        [
            [
                "Separate catalog and runs",
                "Reuse intent across vehicles",
                "Runs bind snapshots later",
            ],
            [
                "Snapshot suite cases",
                "Preserve reviewed composition",
                "Definition edits require a new version",
            ],
            ["Forward-only status", "Avoid accidental reuse", "Archived content remains auditable"],
            ["Bound JSON and lists", "Protect shared resources", "Large data belongs in artifacts"],
            [
                "Materialize cases",
                "Make execution state queryable",
                "Each run owns immutable case identity",
            ],
            [
                "Aggregate in transaction",
                "Keep progress consistent",
                "Case and run update together",
            ],
            [
                "Snapshot when scheduled",
                "Protect delayed execution intent",
                "Dispatch ignores later catalog drift",
            ],
            [
                "Keep native mutation in domains",
                "Preserve simulator ownership",
                "The API remains orchestration-only",
            ],
            [
                "Require recovery intent",
                "Make safe restoration testable",
                "Every step has a bounded recovery plan",
            ],
            [
                "Allowlist mutation operators",
                "Exclude arbitrary execution",
                "Native adapters require isolated execution",
            ],
            [
                "Derive coverage state",
                "Make gaps deterministic",
                "Dashboard consumers use stable totals",
            ],
            [
                "One report per test run",
                "Prevent contradictory final evidence",
                "Exact replay is idempotent",
            ],
            [
                "Validate UI version",
                "Reject stale displayed results",
                "Persistent run remains authoritative",
            ],
        ],
        [1.65, 2.75, 2.5],
    )

    doc.add_page_break()
    doc.add_heading("16 Verification Catalogue", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            ["TF-T-001 to 004", "Definition contracts", "Validate bounds, evidence, and replay"],
            [
                "TF-T-005 to 008",
                "Suite composition",
                "Validate references, snapshots, and evidence",
            ],
            ["TF-T-009 to 010", "Lifecycle", "Validate state and optimistic concurrency"],
            ["TF-T-011 to 012", "RBAC and API", "Validate permissions and safe pagination"],
            ["TF-T-013 to 015", "Persistence", "Validate migration, uniqueness, and atomicity"],
            ["TF-T-016 to 019", "Execution binding", "Validate snapshots, cases, and replay"],
            [
                "TF-T-020 to 026",
                "Case results",
                "Validate bounds, aggregation, APIs, and migration",
            ],
            [
                "TF-T-027 to 034",
                "Scheduled selection",
                "Validate policies, snapshots, dispatch, and migration",
            ],
            [
                "TF-T-035 to 048",
                "Fault campaigns",
                "Validate safety, lifecycle, aggregation, evidence, API, and migration",
            ],
            [
                "TF-T-049 to 059",
                "Mutation and coverage",
                "Validate scoring, traceability, safety, RBAC, integration, and migration",
            ],
            [
                "TF-T-060 to 068",
                "Cross-platform reports",
                "Validate correlation, outcome, replay, safety, RBAC, and integration",
            ],
            [
                "TF-T-069 to 073",
                "Performance and stress",
                "Validate resource bounds, thresholds, baselines, evidence, and migration",
            ],
        ],
        [1.55, 2.45, 2.9],
    )

    doc.add_page_break()
    doc.add_heading("17 Risks and Controls", level=1)
    add_table(
        doc,
        ["Risk", "Control", "Next action"],
        [
            ["Ambiguous sequence", "Unique explicit case order", "Executor follows snapshot order"],
            ["Catalog drift", "Suite and definition snapshots", "Run remains reproducible"],
            [
                "Unbounded results",
                "Duration and collection limits",
                "Large evidence stays in artifacts",
            ],
            [
                "Unsafe activation",
                "Draft review lifecycle and RBAC",
                "Add approval policy if required",
            ],
            ["Large evidence", "Metadata bounds", "Store binary evidence as artifacts"],
            [
                "Catalog changes while waiting",
                "Immutable job selection snapshot",
                "Dispatch only persisted cases",
            ],
            [
                "Unsafe injection scope",
                "Allowlist, blast radius, and time budget",
                "Keep fleet scope excluded",
            ],
            [
                "Vehicle state not restored",
                "Required recovery and verification",
                "Reports retain recovery evidence",
            ],
            [
                "Resource-heavy early benchmarks",
                "Enforce profile budgets and hosted CI",
                "Keep local tests contract-only and GPU-free",
            ],
            [
                "Unsafe mutation execution",
                "Structured allowlist and no generic runner",
                "Require isolated native adapters",
            ],
            [
                "False coverage confidence",
                "Separate test links from evidence links",
                "Track partial and gap totals",
            ],
            [
                "Stale cockpit evidence",
                "Match persisted run status and version",
                "Reject inconsistent reports",
            ],
            [
                "Cross-vehicle correlation",
                "Validate every reference before commit",
                "Keep one authoritative vehicle context",
            ],
        ],
        [1.5, 2.55, 2.85],
    )

    doc.add_heading("18 Study Exercises", level=1)
    doc.add_paragraph(
        "Create one ADAS definition and identify its invariants. Then run two required cases, "
        "explain the aggregate result, and compare an exact retry after suite archival with a new "
        "run request for the archived suite. Schedule the same suite, archive it, and explain why "
        "the stored job can still dispatch reproducibly. Finally, design one battery-temperature "
        "fault with recovery verification and explain why the campaign contract must not expose "
        "an arbitrary shell or adapter command. Create two mutants, calculate a fifty percent "
        "score, and explain why a requirement with a test but no evidence is only partially "
        "covered. Build a final report and identify which reference would be rejected if it came "
        "from another gateway or displayed an older run version."
    )

    doc.add_heading("19 Volume Completion", level=1)
    doc.add_paragraph(
        "The complete Volume VIII baseline is implemented. Performance and stress evidence uses "
        "bounded profiles, deterministic thresholds, and comparable historical baselines. "
        "Meaningful "
        "loads remain isolated from the API and prefer hosted CI to protect workstation resources."
    )
    doc.core_properties.title = "ATEP Volume VIII Test Framework Engineering Workbook"
    doc.core_properties.subject = (
        "Complete functional test framework with cross-platform automation reporting"
    )
    doc.core_properties.author = "ATEP Engineering"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
