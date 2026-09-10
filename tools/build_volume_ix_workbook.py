from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("docs/ATEP_Volume_IX_AI_Test_Engineer_Engineering_Workbook.docx")
BLUE = "1F4E78"
PALE = "EAF1F8"


def shade(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    item = doc.add_table(rows=1, cols=len(headers))
    item.style = "Table Grid"
    header_properties = item.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    for index, header in enumerate(headers):
        cell = item.rows[0].cells[index]
        cell.text = header
        shade(cell, BLUE)
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = item.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[index], PALE)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].font.name = "Aptos"
    title_properties = styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ATEP Volume IX AI Test Engineer Engineering Workbook")
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Version 0.5.0  Explainable Root Cause and Risk").bold = True
    doc.add_paragraph(
        "This workbook records the governed AI boundary and deterministic analysis worker for "
        "ATEP. IX-1 captures analysis intent, IX-2 produces structured advisory results, and "
        "IX-3 converts bounded logs into sanitized timelines and anomaly evidence, while IX-4 "
        "adds reviewed test suggestions, and IX-5 adds explainable risk evidence without a "
        "paid service, model download, network call, or GPU."
    )
    table(
        doc,
        ["Field", "Value"],
        [
            ["Status", "IX-1 through IX-5 implemented"],
            ["Cost", "Local first and no paid dependency"],
            ["Migrations", "0057 foundation through 0062 root cause and risk"],
            ["Next", "IX-6 grounded internal chat"],
        ],
        [1.6, 5.2],
    )
    doc.add_heading("1 Scope and Architecture", level=1)
    doc.add_paragraph(
        "FastAPI accepts bounded provider-neutral requests. A domain service applies idempotency "
        "and data-classification policy. PostgreSQL preserves queued intent, while RBAC, audit, "
        "and outbox evidence provide a traceable boundary. A synchronous application worker uses "
        "versioned local rules behind the same interface reserved for future adapters."
    )
    table(
        doc,
        ["Layer", "Responsibility"],
        [
            ["API", "Validation, pagination, filtering, and RBAC"],
            ["Domain service", "Policy, idempotency, audit, and outbox"],
            ["PostgreSQL", "Immutable request intent, attempts, results, and lifecycle"],
            ["Local worker", "Deterministic versioned rules and safe failure handling"],
            ["Adapter boundary", "Common interface with unknown providers disabled"],
            ["Log intelligence", "Sanitization, timeline, clustering, anomalies, and explanation"],
            ["Test generation", "Requirement-aware draft, human review, and catalog promotion"],
            ["Root cause and risk", "Evidence ranking, risk score, and prediction evaluation"],
        ],
        [1.7, 5.1],
    )
    doc.add_heading("2 Contract and Governance", level=1)
    table(
        doc,
        ["Control", "Rule", "Purpose"],
        [
            [
                "Provider policy",
                "Local only by default",
                "No mandatory external cost or data egress",
            ],
            ["Classification", "Restricted requires local only", "Prevent unsafe transmission"],
            ["Context", "16384 encoded bytes", "Bound storage and processing"],
            ["Evidence", "50 unique references", "Reference artifacts without copying them"],
            ["Authority", "AI remains advisory", "Prevent direct vehicle or test mutation"],
        ],
        [1.55, 2.25, 3.0],
    )
    doc.add_heading("3 Tasks and Subjects", level=1)
    table(
        doc,
        ["Dimension", "Supported values"],
        [
            [
                "Tasks",
                "Log analysis, failure explanation, test suggestion, root cause, risk analysis",
            ],
            ["Subjects", "Test run, automation report, fault execution, mutation execution"],
            ["Lifecycle", "Queued, running, succeeded, or failed with three attempts maximum"],
        ],
        [1.55, 5.25],
    )
    doc.add_heading("4 Deterministic Worker", level=1)
    doc.add_paragraph(
        "The local-rules-v1 adapter evaluates only structured facts supplied in the governed "
        "request. It identifies DTC presence, positive failure counts, and failed performance "
        "thresholds. A no-match result states that no deterministic signal was found instead of "
        "inventing a conclusion."
    )
    table(
        doc,
        ["Control", "Bound", "Engineering purpose"],
        [
            ["Attempts", "Three per request", "Bound retries and resource exposure"],
            ["Execution identity", "Idempotent", "Make worker delivery retry safe"],
            ["Evidence", "Request references only", "Keep conclusions reviewable"],
            ["Failure", "Stable adapter_failure", "Avoid leaking internal exception details"],
            ["Authority", "Advisory output", "Prevent direct state mutation"],
        ],
        [1.55, 2.05, 3.2],
    )
    doc.add_heading("5 Provider Policy and Cost", level=1)
    doc.add_paragraph(
        "Only local-rules is enabled. Unknown providers are rejected before execution. Future "
        "external adapters must satisfy the request provider policy and data classification; "
        "restricted data can never leave the local boundary. IX-2 requires no API key, model, "
        "cloud account, network access, paid service, or GPU."
    )
    doc.add_heading("6 Log Intelligence", level=1)
    doc.add_paragraph(
        "IX-3 accepts up to 500 timestamped lines, sanitizes sensitive values before persistence, "
        "orders timezone-aware events, and clusters messages after replacing volatile identifiers "
        "and numbers. Severe levels and repeated ten-second bursts become deterministic signals."
    )
    table(
        doc,
        ["Control", "Bound", "Engineering purpose"],
        [
            ["Batch", "500 lines and 256000 bytes", "Bound memory and CPU exposure"],
            ["Timeline", "Seven days", "Prevent misleading unbounded correlation"],
            ["Sanitization", "Credentials, email, and VIN", "Reduce sensitive-data persistence"],
            ["Clustering", "Normalized message fingerprint", "Group changing IDs and values"],
            ["Explanation", "Line numbers and evidence refs", "Keep findings reviewable"],
        ],
        [1.5, 2.15, 3.15],
    )
    doc.add_paragraph(
        "Unsupported lines are counted without preserving their content. Explanations describe "
        "observed structure, severity, and proximity and explicitly avoid claiming causality."
    )
    doc.add_heading("7 Governed Test Generation", level=1)
    doc.add_paragraph(
        "IX-4 creates deterministic catalog-compatible drafts from bounded requirement and "
        "evidence "
        "references. Review and promotion are separate versioned actions. Rejection preserves the "
        "reviewer and reason; approval alone does not create or activate a catalog resource."
    )
    table(
        doc,
        ["Control", "Rule", "Engineering purpose"],
        [
            ["Requirements", "At least one bounded reference", "Preserve traceability"],
            ["Review", "Explicit approval or rejection", "Retain human authority"],
            ["Concurrency", "Expected version", "Reject stale decisions"],
            ["Promotion", "Approved suggestions only", "Prevent draft bypass"],
            ["Catalog state", "Definition remains draft", "Prevent automatic execution"],
        ],
        [1.5, 2.15, 3.15],
    )
    doc.add_heading("8 Root Cause and Risk", level=1)
    doc.add_paragraph(
        "IX-5 ranks bounded signals that cite only evidence governed by the request. A transparent "
        "formula combines the highest hypothesis support, severity impact, and inverse "
        "detectability. The result prioritizes investigation and never asserts causality."
    )
    table(
        doc,
        ["Control", "Rule", "Engineering purpose"],
        [
            ["Signals", "1 to 50 unique codes", "Bound processing and ambiguity"],
            ["Evidence", "Request references only", "Prevent unsupported conclusions"],
            ["Risk", "Fixed 0 to 100 formula", "Make prioritization reproducible"],
            ["Outcome", "Versioned observed evidence", "Measure prediction quality"],
            ["Metrics", "Accuracy and Brier score", "Compare historical performance"],
        ],
        [1.5, 2.15, 3.15],
    )
    doc.add_heading("9 Verification Catalogue", level=1)
    table(
        doc,
        ["ID", "Objective"],
        [
            ["AI-T-001", "Validate schema and resource bounds"],
            ["AI-T-002", "Verify local-only defaults"],
            ["AI-T-003", "Reject restricted external processing"],
            ["AI-T-004", "Verify idempotency and conflict"],
            ["AI-T-005", "Verify minimized audit and event evidence"],
            ["AI-T-006", "Verify RBAC and OpenAPI contracts"],
            ["AI-T-007", "Apply migration through hosted Docker integration"],
            ["AI-T-008 to 009", "Verify deterministic rules and complete lifecycle"],
            ["AI-T-010 to 011", "Verify retries, replay, adapter policy, and conflicts"],
            ["AI-T-012 to 013", "Verify cited output, minimized evidence, API, and RBAC"],
            ["AI-T-014", "Apply migration 0059 through hosted Docker integration"],
            ["AI-T-015 to 017", "Verify parsing, sanitization, ordering, and clustering"],
            ["AI-T-018 to 019", "Verify anomaly evidence and grounded explanations"],
            ["AI-T-020 to 021", "Verify bounds, replay, APIs, RBAC, and minimized events"],
            ["AI-T-022", "Apply migration 0060 through hosted Docker integration"],
            ["AI-T-023 to 025", "Verify generation bounds, replay, review, and rejection evidence"],
            ["AI-T-026 to 029", "Verify promotion gate, draft catalog state, APIs, and RBAC"],
            ["AI-T-030", "Apply migration 0061 through hosted Docker integration"],
            ["AI-T-031 to 034", "Verify signal policy, ranking, risk formula, and limitations"],
            ["AI-T-035 to 037", "Verify evaluation, Brier score, metrics, APIs, and RBAC"],
            ["AI-T-038", "Apply migration 0062 through hosted Docker integration"],
        ],
        [1.4, 5.4],
    )
    doc.add_heading("10 Risks and Next Development", level=1)
    table(
        doc,
        ["Risk", "Control"],
        [
            ["Sensitive data egress", "Classification and provider policy gate"],
            ["Hallucinated authority", "Advisory-only boundary and future citations"],
            ["Unexpected cost", "No provider call and local-only default"],
            ["Resource consumption", "No model or GPU in IX-1 through IX-5"],
            ["Prompt leakage", "Minimized audit and outbox metadata"],
            ["Unbounded retries", "Three immutable attempts maximum"],
            ["Worker exception leakage", "Stable error code without exception text"],
            ["Sensitive log values", "Sanitize before persistence and omit rejected text"],
            ["False causal claim", "Report signals, citations, and explicit limitations"],
            ["Unsafe generated test", "Human review and inactive catalog promotion"],
            ["False root cause", "Rank associations with explicit causal limitations"],
            ["Misread risk score", "Document fixed formula and non-probability meaning"],
        ],
        [2.2, 4.6],
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
