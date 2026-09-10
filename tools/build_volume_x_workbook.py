from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path("docs/ATEP_Volume_X_Dashboard_Engineering_Workbook.docx")
BLUE = "1F4E78"
PALE = "EAF1F8"


def shade(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(marker)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, PALE)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ATEP Volume X Dashboard Engineering Workbook")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Version 0.1.0  Dashboard Foundation").bold = True
    doc.add_paragraph(
        "This workbook records X-1, the first governed dashboard read model. It consolidates "
        "test execution, quality, automation, mutation, and AI evidence into a stable FastAPI "
        "contract without duplicating source data or introducing a paid service, model, or GPU."
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Status", "X-1 complete"],
            ["Contract", "dashboard-overview-v1"],
            ["Endpoint", "GET /api/v1/dashboard/overview"],
            ["Authorization", "dashboard:read"],
            ["Next", "X-2 test quality trends and failure drill-down"],
        ],
        [1.65, 5.15],
    )

    doc.add_heading("1 Scope and Architecture", level=1)
    doc.add_paragraph(
        "The dashboard service is a read-only projection over authoritative ATEP tables. FastAPI "
        "validates bounded query parameters, SQLAlchemy executes aggregate queries in PostgreSQL, "
        "and Pydantic publishes a deterministic client contract. No dashboard table, event, audit "
        "mutation, or frontend framework is required for X-1."
    )
    add_table(
        doc,
        ["Layer", "Responsibility"],
        [
            ["API", "Authentication, dashboard RBAC, window limits, and response contract"],
            ["Read service", "Aggregate counts, percentages, score averages, and recent cards"],
            [
                "Source domains",
                "Remain authoritative for tests, quality, automation, and AI evidence",
            ],
            ["Client contract", "Supply KPI, chart, heatmap, and evidence-list data"],
        ],
        [1.6, 5.2],
    )

    doc.add_heading("2 Contract and KPI Semantics", level=1)
    add_table(
        doc,
        ["Area", "Measures", "Time semantics"],
        [
            [
                "Test execution",
                "Runs, active runs, case totals, passed, failed, pass rate",
                "Selected window",
            ],
            [
                "Requirements",
                "Total, covered, coverage rate, status distribution",
                "Current snapshot",
            ],
            ["Mutation", "Execution count and average mutation score", "Selected window"],
            ["Automation", "Report count and outcome distribution", "Selected window"],
            [
                "AI evidence",
                "Dashboard projection count, severity, and recent cited cards",
                "Selected window",
            ],
        ],
        [1.35, 3.5, 1.95],
    )
    doc.add_paragraph(
        "Undefined rates are returned as null when no denominator exists. Counts and distributions "
        "are sorted deterministically; mutation scores and percentages are rounded to two decimals."
    )

    doc.add_heading("3 Security Privacy and Bounds", level=1)
    add_table(
        doc,
        ["Control", "Rule", "Purpose"],
        [
            ["RBAC", "dashboard:read", "Separate dashboard visibility from source management"],
            ["Window", "1 to 720 hours", "Bound database work and interpretation range"],
            ["Evidence", "1 to 50 cards", "Bound payload, memory, and explanatory content"],
            [
                "Minimization",
                "Governed projection fields only",
                "Exclude prompts, raw logs, and context",
            ],
            ["Authority", "Read only", "Prevent dashboard requests from changing operations"],
        ],
        [1.35, 2.25, 3.2],
    )

    doc.add_heading("4 Requirements Traceability", level=1)
    add_table(
        doc,
        ["Requirement group", "Implemented evidence"],
        [
            [
                "DASH-F-001 to 003",
                "Authenticated endpoint, execution aggregation, and safe pass rate",
            ],
            ["DASH-F-004 to 006", "Coverage, quality sources, and cited minimized AI cards"],
            ["DASH-NF-001 to 003", "Dedicated permission and strict window and payload limits"],
            [
                "DASH-NF-004 to 006",
                "Read-only design, deterministic output, and local resource profile",
            ],
        ],
        [2.0, 4.8],
    )

    doc.add_heading("5 Test Catalogue", level=1)
    add_table(
        doc,
        ["Test", "Objective"],
        [
            [
                "KPI calculation",
                "Verify totals, active states, rates, averages, and deterministic rounding",
            ],
            ["Empty state", "Prevent missing evidence from appearing as a zero-percent result"],
            ["Evidence mapping", "Preserve governed citations and bounded presentation fields"],
            ["OpenAPI bounds", "Publish safe defaults and maximum query limits"],
            ["RBAC", "Reject authenticated users without dashboard:read with HTTP 403"],
            [
                "Integration",
                "Exercise the endpoint against PostgreSQL and the deployed permission catalogue",
            ],
            ["Regression", "Protect every previously delivered domain and integration contract"],
            ["Static quality", "Keep formatting and type analysis clean with Ruff and mypy"],
        ],
        [2.0, 4.8],
    )

    doc.add_heading("6 Risks Decisions and Technical Debt", level=1)
    add_table(
        doc,
        ["Topic", "Current decision", "Future trigger"],
        [
            [
                "Query cost",
                "Bound live aggregation",
                "Materialize only after measured latency requires it",
            ],
            ["Historical trends", "Not included in X-1", "Add bounded buckets in X-2"],
            [
                "Certification",
                "Explicitly not inferred",
                "Add evidence views, not certification claims",
            ],
            [
                "Frontend",
                "Contract first",
                "Choose UI stack after stable workflows and accessibility needs",
            ],
            ["Live data", "Request response only", "Add controlled streaming in a later increment"],
        ],
        [1.35, 2.7, 2.75],
    )

    doc.add_heading("7 Cost and Resource Profile", level=1)
    doc.add_paragraph(
        "X-1 reuses the local FastAPI and PostgreSQL stack. It requires no AWS account, paid API, "
        "external AI provider, model download, or GPU. The automated tests use small deterministic "
        "fixtures and bounded database checks, making the increment appropriate for the current "
        "Windows and Docker development environment."
    )

    doc.add_heading("8 Next Planned Increment", level=1)
    doc.add_paragraph(
        "X-2 will add time-bucketed test quality trends and a bounded failure drill-down contract. "
        "It will reference existing test runs, case results, and artifacts without copying raw "
        "evidence into the dashboard domain."
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
