from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ATEP_Volume_VII_ADAS_Engineering_Workbook.docx"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, (cell, value) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.width = Inches(widths[index])
        shade(cell, "17365D")
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values, strict=True)):
            cell.width = Inches(widths[index])
            if row_index % 2:
                shade(cell, "EAF0F7")
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def bullets(doc: Document, values: list[str]) -> None:
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 26), ("Heading 1", 17), ("Heading 2", 13)]:
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ATEP Volume VII ADAS Engineering Workbook")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Version 0.6.0   VII 1 through VII 6 Implemented").bold = True
    doc.add_paragraph(
        "This workbook records the first engineering baseline for the ATEP ADAS volume. "
        "The world model provides deterministic sensor independent ground truth, so later "
        "perception and planning components can be tested against an authoritative scene."
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Status", "VII-1 through VII-6 implemented and verified"],
            ["Technology", "FastAPI, PostgreSQL, SQLAlchemy, Alembic, Pydantic"],
            ["Cost", "Local first with no paid cloud or AI dependency"],
            ["Next", "VII-7 cross-platform integration"],
        ],
        [1.5, 5.4],
    )

    doc.add_heading("1 Scope and Outcome", level=1)
    doc.add_paragraph(
        "VII-1 establishes the scene truth used by every future ADAS test. It includes an "
        "ENU coordinate frame, bounded road and lane geometry, exactly one ego vehicle, "
        "other traffic actors, deterministic logical time, persistence, RBAC, audit records, "
        "and transactional integration events."
    )
    bullets(
        doc,
        [
            "In scope: creation, lookup, pagination, and constant velocity advancement.",
            "In scope: ego vehicle, vehicle, pedestrian, cyclist, and static obstacle actors.",
            "Delivered next: environment, sensors, perception, planning, alerts, and scenarios.",
            "Deferred: cross-platform test-run, Vehicle Gateway, CarSystemUI, and dashboard links.",
        ],
    )

    doc.add_heading("2 Architecture", level=1)
    doc.add_paragraph(
        "The Android Automotive interface and future simulators continue to communicate "
        "through ATEP APIs. The ADAS module owns world truth. Sensor modules will consume a "
        "scene revision and produce observations without changing that truth."
    )
    add_table(
        doc,
        ["Layer", "Responsibility"],
        [
            ["FastAPI router", "HTTP contracts, RBAC, pagination, and transaction boundaries"],
            ["Domain service", "Deterministic advancement, audit, and outbox evidence"],
            ["Pydantic", "Bounds and invariants for coordinates, geometry, actors, and time"],
            [
                "PostgreSQL",
                "Vehicle scoped scene, revision, logical time, roads, lanes, and actors",
            ],
            ["Alembic", "Reproducible migration 0044 and downgrade"],
        ],
        [1.55, 5.35],
    )

    doc.add_heading("3 Domain Model and Contracts", level=1)
    add_table(
        doc,
        ["Concept", "Invariant"],
        [
            ["Coordinate frame", "East North Up with bounded WGS84 origin"],
            ["Road and lane", "Unique identifiers and at least two centerline points per lane"],
            ["Actor", "Unique identifier with bounded dimensions, pose, heading, and velocity"],
            ["Ego vehicle", "Exactly one per scene"],
            ["Revision", "Starts at one and increments for every accepted advance"],
            ["Logical time", "Starts at zero and changes only through an advance command"],
        ],
        [1.55, 5.35],
    )

    doc.add_heading("4 API Surface", level=1)
    add_table(
        doc,
        ["Method", "Path", "Permission", "Purpose"],
        [
            ["POST", "/vehicles/{vehicle_id}/adas/scenes", "adas:manage", "Create truth"],
            ["GET", "/vehicles/{vehicle_id}/adas/scenes", "adas:read", "List scenes"],
            ["GET", "/vehicles/{vehicle_id}/adas/scenes/{scene_id}", "adas:read", "Read scene"],
            [
                "POST",
                "/vehicles/{vehicle_id}/adas/scenes/{scene_id}/advance",
                "adas:manage",
                "Advance time",
            ],
            [
                "PATCH",
                "/vehicles/{vehicle_id}/adas/scenes/{scene_id}/context",
                "adas:manage",
                "Update environment and traffic controls",
            ],
        ],
        [0.65, 3.35, 1.2, 1.25],
    )

    doc.add_heading("5 Engineering Decisions", level=1)
    add_table(
        doc,
        ["Decision", "Rationale", "Consequence"],
        [
            [
                "Separate truth from detections",
                "Allows objective perception scoring",
                "Sensors reference a scene revision",
            ],
            [
                "Use ENU locally",
                "Supports simple metric calculations",
                "Large area projection is deferred",
            ],
            [
                "Use logical time",
                "Makes tests independent of wall clock",
                "Time changes only through commands",
            ],
            [
                "Use expected revision",
                "Prevents silent concurrent overwrite",
                "Stale writes return stable HTTP 409",
            ],
            [
                "Store bounded JSON aggregates",
                "Supports heterogeneous actors",
                "Schema evolution needs contract discipline",
            ],
        ],
        [1.7, 2.65, 2.45],
    )

    doc.add_heading("6 Verification Catalogue", level=1)
    definitions = [
        ("Valid ENU scene", "Accept complete ground truth"),
        ("Single ego rule", "Reject missing or multiple ego actors"),
        ("Identifier uniqueness", "Prevent ambiguous roads, lanes, and actors"),
        ("Physical bounds", "Reject unsafe inputs"),
        ("Deterministic advance", "Verify positions and logical time"),
        ("Stale revision", "Prevent mutation on conflict"),
        ("Creation evidence", "Prove atomic audit and event evidence"),
        ("Advance evidence", "Keep event evidence minimized"),
        ("RBAC", "Prove read and manage protection"),
        ("Pagination", "Protect collection resources"),
        ("Migration", "Prove upgrade and downgrade"),
        ("Isolated replay", "Prove identical results from identical inputs"),
    ]
    rows = [
        [f"ADAS-T-{index:03d}", name, objective]
        for index, (name, objective) in enumerate(definitions, 1)
    ]
    add_table(doc, ["ID", "Test", "Objective"], rows, [1.05, 2.15, 3.65])

    doc.add_heading("7 Evidence and Quality Gates", level=1)
    bullets(
        doc,
        [
            "Focused ADAS and API contract tests pass in one process.",
            "Ruff formatting and static checks pass for the new module.",
            "mypy passes for the module and application composition.",
            "Migrations 0044 through 0049 form one linear ADAS migration history.",
            "The implementation uses no GPU and no paid external service.",
        ],
    )
    doc.add_heading("8 Risks and Controls", level=1)
    add_table(
        doc,
        ["Risk", "Control", "Future action"],
        [
            [
                "Unbounded scenes",
                "Limits for roads, lanes, actors, and points",
                "Add payload metrics",
            ],
            ["Floating point drift", "Round positions to six decimals", "Evaluate fixed point"],
            [
                "Truth coupled to sensors",
                "Separate scenes from observation models",
                "Link observations to revisions",
            ],
            ["Concurrent mutation", "Expected revision conflict", "Add idempotent command IDs"],
            [
                "Simplified motion",
                "Constant velocity plus deterministic waypoint trajectories",
                "Add acceleration profiles after VII-3",
            ],
        ],
        [1.65, 2.8, 2.1],
    )
    doc.add_page_break()
    doc.add_heading("9 Environment Traffic and Trajectories", level=1)
    doc.add_paragraph(
        "VII-2 extends scene ground truth with bounded weather, precipitation, visibility, "
        "illumination, temperature, wind, and road friction. Traffic controls reference known "
        "lanes and carry type-specific state. Actor trajectories use ordered absolute logical-time "
        "waypoints and deterministic linear interpolation."
    )
    add_table(
        doc,
        ["Contract", "Control", "Test objective"],
        [
            [
                "Weather",
                "Cross-field precipitation and fog validation",
                "Reject inconsistent conditions",
            ],
            [
                "Road friction",
                "Coefficient from 0.05 through 1.5",
                "Represent low and high grip safely",
            ],
            [
                "Traffic control",
                "Unique ID and references to existing lanes",
                "Prevent ambiguous map truth",
            ],
            [
                "Trajectory",
                "Starts at zero with strictly increasing times",
                "Guarantee a valid timeline",
            ],
            [
                "Context update",
                "Expected revision and atomic evidence",
                "Prevent lost concurrent updates",
            ],
        ],
        [1.55, 2.75, 2.55],
    )
    doc.add_heading("10 VII 2 Verification Addendum", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            [
                "ADAS-T-013 to 014",
                "Weather consistency",
                "Validate precipitation and visibility rules",
            ],
            [
                "ADAS-T-015 to 017",
                "Traffic control contracts",
                "Validate lane references, fields, and identity",
            ],
            [
                "ADAS-T-018 to 020",
                "Trajectory timeline",
                "Validate ordering, interpolation, and terminal motion",
            ],
            [
                "ADAS-T-021 to 022",
                "Context mutation",
                "Validate concurrency, audit, and event evidence",
            ],
        ],
        [1.55, 2.55, 2.75],
    )
    doc.add_heading("11 Deterministic Sensor Simulation", level=1)
    doc.add_paragraph(
        "VII-3 adds persisted camera, radar, and LiDAR configurations and observations. "
        "Every observation identifies its scene revision, captured logical time, latency-adjusted "
        "observation time, and deterministic seed while leaving scene ground truth unchanged."
    )
    add_table(
        doc,
        ["Mechanism", "Implementation", "Verification objective"],
        [
            [
                "Geometry",
                "Mount pose, yaw, range, and horizontal field of view",
                "Exclude actors outside the sensor envelope",
            ],
            [
                "Occlusion",
                "Near-to-far angular coverage",
                "Hide aligned actors behind nearer actors",
            ],
            [
                "Noise",
                "SHA-256 seed-derived bounded position offsets",
                "Reproduce identical detections",
            ],
            [
                "Environment",
                "Visibility range and sensor-sensitive confidence",
                "Represent rain, fog, snow, and low light",
            ],
            [
                "Evidence",
                "Persisted observation plus minimized audit and outbox event",
                "Retain traceability without event bloat",
            ],
        ],
        [1.4, 2.9, 2.55],
    )
    doc.add_heading("12 VII 3 Verification Addendum", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            [
                "ADAS-T-023 to 025",
                "Bounds, FOV, and range",
                "Validate configuration and spatial filtering",
            ],
            [
                "ADAS-T-026 to 029",
                "Occlusion, noise, and environment",
                "Validate deterministic sensor effects",
            ],
            [
                "ADAS-T-030 to 032",
                "Revision, time, and identity",
                "Validate stable persisted evidence",
            ],
            [
                "ADAS-T-033 to 034",
                "RBAC, audit, and events",
                "Validate access and atomic traceability",
            ],
        ],
        [1.55, 2.55, 2.75],
    )
    doc.add_heading("13 Perception and Ground Truth Scoring", level=1)
    doc.add_paragraph(
        "VII-4 persists perception outputs for objects, pedestrians, lanes, signs, and signals. "
        "Predictions reference stable truth identifiers and are matched one-to-one against actors "
        "visible in the sensor observation plus lane and traffic-control truth from the unchanged "
        "scene revision."
    )
    add_table(
        doc,
        ["Metric", "Definition", "Engineering control"],
        [
            ["Precision", "TP divided by TP plus FP", "Measures prediction correctness"],
            ["Recall", "TP divided by TP plus FN", "Measures ground-truth coverage"],
            ["F1", "Harmonic mean of precision and recall", "Balances both failure modes"],
            ["Revision gate", "Observation revision equals current scene", "Rejects stale truth"],
            ["One-to-one match", "Each truth target matches once", "Prevents score inflation"],
        ],
        [1.35, 2.8, 2.7],
    )
    doc.add_heading("14 VII 4 Verification Addendum", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            ["ADAS-T-035 to 038", "Contracts and truth", "Validate all perception target types"],
            ["ADAS-T-039 to 041", "Deterministic scoring", "Validate exact one-to-one metrics"],
            ["ADAS-T-042", "Revision gate", "Reject comparison against changed truth"],
            ["ADAS-T-043 to 044", "RBAC and evidence", "Validate access and traceability"],
        ],
        [1.55, 2.55, 2.75],
    )
    doc.add_page_break()
    doc.add_heading("15 Planning and Alerts", level=1)
    doc.add_paragraph(
        "VII-5 evaluates perceived actors, the declared ego lane, and perceived signal state. "
        "Relative forward motion produces collision TTC and lead-vehicle distance, while nearest-"
        "segment geometry produces lane-center offset. Ordered risks select one deterministic "
        "maneuver and a stable alert sequence."
    )
    add_table(
        doc,
        ["Decision input", "Evaluation", "Possible outcome"],
        [
            ["Perceived actors", "Forward path, relative speed, TTC", "Brake or emergency brake"],
            ["Lead vehicle", "Minimum following distance", "Unsafe following warning"],
            ["Ego lane", "Centerline offset and usable envelope", "Lane centering"],
            ["Perceived signal", "Red classification", "Stop and critical alert"],
            ["Concurrent risks", "Fixed maneuver priority", "Reproducible decision"],
        ],
        [1.55, 2.8, 2.3],
    )
    doc.add_heading("16 VII 5 Verification Addendum", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            ["ADAS-T-045 to 047", "Thresholds and motion", "Validate bounds, distance, and TTC"],
            ["ADAS-T-048 to 052", "Alerts and maneuvers", "Validate risk decisions and priority"],
            ["ADAS-T-053", "Truth integrity", "Reject unknown lanes and stale revisions"],
            ["ADAS-T-054 to 055", "RBAC and evidence", "Validate access and traceability"],
        ],
        [1.55, 2.55, 2.75],
    )
    doc.add_page_break()
    doc.add_heading("17 ADAS Test Scenarios", level=1)
    doc.add_paragraph(
        "VII-6 executes bounded, repeatable engineering scenarios over an immutable scene and "
        "persisted perception result. The catalogue is inspired by common consumer safety "
        "assessment concerns but does not claim official NCAP certification or homologation."
    )
    add_table(
        doc,
        ["Scenario family", "Primary behavior", "Evidence"],
        [
            ["Car to car AEB", "Lead vehicle collision response", "TTC alert and braking"],
            ["Pedestrian AEB", "Pedestrian collision response", "Detection and maneuver"],
            ["Lane support", "Lane envelope recovery", "Departure alert and centering"],
            ["Traffic signal", "Red signal compliance", "Critical alert and stop"],
        ],
        [1.55, 2.8, 2.3],
    )
    doc.add_heading("18 Fault Regression and Coverage", level=1)
    add_table(
        doc,
        ["Mechanism", "Engineering control", "Purpose"],
        [
            [
                "Fault injection",
                "Drop or misclassify up to twenty unique predictions",
                "Exercise degraded perception safely",
            ],
            [
                "Assertions",
                "Expected maneuver, required alerts, and minimum F1",
                "Produce explicit pass or fail evidence",
            ],
            [
                "Coverage",
                "Scenario, target, alert, maneuver, fault, and assertion dimensions",
                "Show what each execution exercised",
            ],
            [
                "Fingerprint",
                "Canonical deterministic SHA-256 evidence",
                "Compare regressions across execution identities",
            ],
            [
                "Replay",
                "Request hash and stable conflict",
                "Prevent duplicate or ambiguous evidence",
            ],
        ],
        [1.4, 3.0, 2.25],
    )
    doc.add_heading("19 VII 6 Verification Addendum", level=1)
    add_table(
        doc,
        ["ID range", "Coverage", "Objective"],
        [
            ["ADAS-T-056 to 060", "Contracts and faults", "Validate bounded fault behavior"],
            ["ADAS-T-061 to 064", "Evidence and replay", "Validate results and determinism"],
            ["ADAS-T-065 to 068", "API and persistence", "Validate access and traceability"],
        ],
        [1.55, 2.55, 2.75],
    )
    doc.add_heading("20 Next Development", level=1)
    doc.add_paragraph(
        "VII-7 will connect ADAS scenario evidence to ATEP test runs, Vehicle Gateway, "
        "CarSystemUI presentation, and real-time dashboard streams."
    )
    doc.add_heading("21 Study Exercises", level=1)
    bullets(
        doc,
        [
            "Repeat an advance with a stale revision and explain the stable conflict.",
            "Submit one duplicate and one misclassified prediction and calculate the F1 score.",
            "Compare braking decisions at TTC values above and below the emergency threshold.",
            "Drop the lead-vehicle prediction, compare fingerprints, and explain why these "
            "scenarios are not official certification evidence.",
        ],
    )
    doc.core_properties.title = "ATEP Volume VII ADAS Engineering Workbook"
    doc.core_properties.subject = "Deterministic ADAS world model engineering evidence"
    doc.core_properties.author = "ATEP Engineering"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
